"""Excel 扫描器 V2 测试"""
import pytest
from pathlib import Path

from openpyxl import Workbook

from src.finance_mask.scanner import ExcelScannerV2, SheetType, LayoutViewClassifier
from src.finance_mask.scanner.header_finder import find_header_row


class TestExcelScannerV2:
    """Excel 扫描器 V2 测试"""

    @pytest.fixture
    def data_workbook(self, tmp_path):
        """创建 Data 类型的工作簿（表头+数据的一维表格）"""
        wb = Workbook()
        ws = wb.active
        ws.title = "利润表"

        # 表头行
        headers = ["项目", "本期金额", "上期金额", "同比变动"]
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)

        # 数据行
        data = [
            ("营业收入", 12345678.90, 11111111.11, "11.1%"),
            ("营业成本", 8765432.10, 7777777.77, "12.7%"),
            ("毛利润", 3580246.80, 3333333.34, "7.4%"),
            ("净利润", 2345678.90, 2222222.22, "5.6%"),
        ]

        for row_idx, (item, current, prev, change) in enumerate(data, start=2):
            ws.cell(row=row_idx, column=1, value=item)
            ws.cell(row=row_idx, column=2, value=current)
            ws.cell(row=row_idx, column=3, value=prev)
            ws.cell(row=row_idx, column=4, value=change)

        filepath = tmp_path / "data_report.xlsx"
        wb.save(str(filepath))
        return filepath

    @pytest.fixture
    def form_workbook(self, tmp_path):
        """创建 Form 类型的工作簿（key-value 排版的表单数据）"""
        wb = Workbook()
        ws = wb.active
        ws.title = "基本信息"

        # 表单数据（key-value 格式）
        form_data = [
            ("公司名称", "阿里巴巴集团"),
            ("统一社会信用代码", "91330000799210495T"),
            ("法定代表人", "张三"),
            ("注册资本", "1000000万元"),
            ("成立日期", "2020-01-01"),
            ("营业收入", "1234567890.12元"),
            ("净利润", "234567890.50元"),
        ]

        for row_idx, (key, value) in enumerate(form_data, start=1):
            ws.cell(row=row_idx, column=1, value=key)
            ws.cell(row=row_idx, column=2, value=value)

        filepath = tmp_path / "form_report.xlsx"
        wb.save(str(filepath))
        return filepath

    def test_find_header_row(self, data_workbook):
        """测试表头行查找"""
        from openpyxl import load_workbook

        wb = load_workbook(str(data_workbook), read_only=True)
        ws = wb.active

        header_start, header_end = find_header_row(ws, candidates=None)
        wb.close()

        # 表头行应该是第 0 行（索引）
        assert header_start == 0
        assert header_end == 0  # 单行表头

    def test_scan_data_sheet(self, data_workbook):
        """测试扫描 Data 类型工作表"""
        scanner = ExcelScannerV2()
        sites = scanner.scan(data_workbook)

        # 应该发现敏感位点
        assert len(sites) > 0
        print(f"发现 {len(sites)} 个敏感位点")

        # 检查是否包含金额位点
        amount_sites = [s for s in sites if s.detected_type.value == "amount"]
        assert len(amount_sites) > 0

    def test_scan_form_sheet(self, form_workbook):
        """测试扫描 Form 类型工作表"""
        scanner = ExcelScannerV2()
        sites = scanner.scan(form_workbook)

        # 应该发现敏感位点
        assert len(sites) > 0
        print(f"发现 {len(sites)} 个敏感位点")

        # 检查是否包含金额位点（Form 表单中的金额）
        amount_sites = [s for s in sites if s.detected_type.value == "amount"]
        assert len(amount_sites) > 0

    def test_sheet_type_classification(self, data_workbook, form_workbook):
        """测试工作表类型分类"""
        classifier = LayoutViewClassifier()

        # 测试 Data 类型工作簿
        classifications = classifier.classify_with_fallback(data_workbook)
        assert len(classifications) > 0
        print(f"Data 工作簿分类: {classifications}")

        # 测试 Form 类型工作簿
        classifications = classifier.classify_with_fallback(form_workbook)
        assert len(classifications) > 0
        print(f"Form 工作簿分类: {classifications}")

    def test_scanner_with_multiple_sheets(self, tmp_path):
        """测试包含多个工作表的扫描"""
        wb = Workbook()

        # 创建 Data 类型工作表（带有多行表头）
        ws1 = wb.active
        ws1.title = "利润表"
        
        # 第一行：大类表头
        ws1.cell(row=1, column=1, value="项目")
        ws1.cell(row=1, column=2, value="金额")
        ws1.cell(row=1, column=3, value="金额")
        
        # 第二行：小类表头
        ws1.cell(row=2, column=1, value="")
        ws1.cell(row=2, column=2, value="本期")
        ws1.cell(row=2, column=3, value="上期")
        
        # 数据行
        ws1.cell(row=3, column=1, value="营业收入")
        ws1.cell(row=3, column=2, value=12345678.90)
        ws1.cell(row=3, column=3, value=11111111.11)

        # 创建 Form 类型工作表
        ws2 = wb.create_sheet("基本信息")
        ws2.cell(row=1, column=1, value="公司名称")
        ws2.cell(row=1, column=2, value="阿里巴巴集团")
        ws2.cell(row=2, column=1, value="营业收入")
        ws2.cell(row=2, column=2, value="1234567890.12元")

        filepath = tmp_path / "multi_sheet.xlsx"
        wb.save(str(filepath))

        # 扫描
        scanner = ExcelScannerV2()
        sites = scanner.scan(filepath)

        # 应该发现多个位点
        assert len(sites) > 0
        print(f"多工作表扫描发现 {len(sites)} 个敏感位点")

        # 检查不同工作表的位点
        sheet_names = set(s.location.sheet for s in sites)
        assert len(sheet_names) >= 1  # 至少有一个工作表

    def test_data_sheet_column_matching(self, data_workbook):
        """测试 Data 类型工作表的列头匹配"""
        scanner = ExcelScannerV2()
        sites = scanner.scan(data_workbook)

        # 检查列头匹配结果
        for site in sites:
            if site.location.column:
                print(f"列 '{site.location.column}' -> 位点 {site.site_id}, 类型: {site.detected_type}")

        # 应该有列头匹配的位点
        column_rule_sites = [s for s in sites if s.discovered_by.value == "column_rule"]
        assert len(column_rule_sites) > 0

    @pytest.fixture
    def multi_header_workbook(self, tmp_path):
        """创建包含多行表头的工作簿"""
        wb = Workbook()
        ws = wb.active
        ws.title = "财务报表"

        # 第一行：大类表头
        ws.cell(row=1, column=1, value="项目")
        ws.cell(row=1, column=2, value="本期")
        ws.cell(row=1, column=3, value="本期")
        ws.cell(row=1, column=4, value="上期")
        ws.cell(row=1, column=5, value="上期")

        # 第二行：小类表头
        ws.cell(row=2, column=1, value="")
        ws.cell(row=2, column=2, value="金额")
        ws.cell(row=2, column=3, value="占比")
        ws.cell(row=2, column=4, value="金额")
        ws.cell(row=2, column=5, value="占比")

        # 数据行
        ws.cell(row=3, column=1, value="营业收入")
        ws.cell(row=3, column=2, value=12345678.90)
        ws.cell(row=3, column=3, value="100%")
        ws.cell(row=3, column=4, value=11111111.11)
        ws.cell(row=3, column=5, value="100%")

        ws.cell(row=4, column=1, value="营业成本")
        ws.cell(row=4, column=2, value=8765432.10)
        ws.cell(row=4, column=3, value="71%")
        ws.cell(row=4, column=4, value=7777777.77)
        ws.cell(row=4, column=5, value="70%")

        filepath = tmp_path / "multi_header.xlsx"
        wb.save(str(filepath))
        return filepath

    def test_multi_header_row(self, multi_header_workbook):
        """测试多行表头查找"""
        from openpyxl import load_workbook

        wb = load_workbook(str(multi_header_workbook), read_only=True)
        ws = wb.active

        header_start, header_end = find_header_row(ws, candidates=None)
        wb.close()

        # 应该找到多行表头
        assert header_start == 0
        assert header_end == 1  # 两行表头

    def test_scan_multi_header_sheet(self, multi_header_workbook):
        """测试扫描多行表头工作表"""
        scanner = ExcelScannerV2()
        sites = scanner.scan(multi_header_workbook)

        # 应该发现敏感位点
        assert len(sites) > 0
        print(f"多行表头工作表发现 {len(sites)} 个敏感位点")

        # 检查列头匹配结果
        for site in sites:
            if site.location.column:
                print(f"列 '{site.location.column}' -> 位点 {site.site_id}, 类型: {site.detected_type}")


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
