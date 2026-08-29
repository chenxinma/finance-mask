"""零宽字符水印解码"""
from __future__ import annotations

import re


class WatermarkDecoder:
    """零宽字符水印解码器"""

    # 零宽字符到 bit 的映射
    CHAR_TO_BIT = {
        "\u200B": "0",  # 零宽空格
        "\u200C": "1",  # 零宽非连接符
        "\u200D": "1",  # 零宽连接符（备选）
        "\uFEFF": "0",  # 零宽不换行空格（备选）
    }

    @classmethod
    def decode(cls, text: str) -> str:
        """
        从文本中提取并解码零宽字符水印
        
        Args:
            text: 包含零宽字符的文本
            
        Returns:
            解码后的 payload
            
        Raises:
            ValueError: 无法解码或解码结果无效
        """
        if not text:
            raise ValueError("空文本无法解码")

        # 提取所有零宽字符
        zero_width_pattern = re.compile(r"[\u200B\u200C\u200D\uFEFF]")
        zero_width_chars = zero_width_pattern.findall(text)

        if not zero_width_chars:
            raise ValueError("未找到零宽字符水印")

        # 转换为二进制字符串
        binary = ""
        for char in zero_width_chars:
            if char in cls.CHAR_TO_BIT:
                binary += cls.CHAR_TO_BIT[char]

        # 检查长度是否为 8 的倍数
        if len(binary) % 8 != 0:
            # 截断到 8 的倍数
            binary = binary[: len(binary) - (len(binary) % 8)]

        if not binary:
            raise ValueError("零宽字符解码失败")

        # 转换为字节
        byte_list = []
        for i in range(0, len(binary), 8):
            byte_str = binary[i : i + 8]
            if len(byte_str) == 8:
                byte_list.append(int(byte_str, 2))

        # UTF-8 解码
        try:
            payload = bytes(byte_list).decode("utf-8")
            return payload
        except UnicodeDecodeError:
            raise ValueError("水印解码失败：UTF-8 解码错误")

    @classmethod
    def has_watermark(cls, text: str) -> bool:
        """
        检测文本是否包含水印
        
        Args:
            text: 要检测的文本
            
        Returns:
            是否包含水印
        """
        if not text:
            return False

        zero_width_pattern = re.compile(r"[\u200B\u200C\u200D\uFEFF]")
        return bool(zero_width_pattern.search(text))

    @classmethod
    def extract_from_excel(cls, filepath: str) -> list[str]:
        """
        从 Excel 文件中提取水印
        
        Args:
            filepath: Excel 文件路径
            
        Returns:
            解码后的 payload 列表
        """
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filepath, read_only=True)
            payloads = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row in ws.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str):
                            if cls.has_watermark(cell.value):
                                try:
                                    payload = cls.decode(cell.value)
                                    payloads.append(payload)
                                except ValueError:
                                    pass

            wb.close()
            return list(set(payloads))  # 去重

        except Exception as e:
            raise ValueError(f"Excel 水印提取失败: {e}")

    @classmethod
    def extract_from_ppt(cls, filepath: str) -> list[str]:
        """
        从 PPT 文件中提取水印
        
        Args:
            filepath: PPT 文件路径
            
        Returns:
            解码后的 payload 列表
        """
        try:
            from pptx import Presentation

            prs = Presentation(filepath)
            payloads = []

            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = paragraph.text
                            if text and cls.has_watermark(text):
                                try:
                                    payload = cls.decode(text)
                                    payloads.append(payload)
                                except ValueError:
                                    pass

            return list(set(payloads))  # 去重

        except Exception as e:
            raise ValueError(f"PPT 水印提取失败: {e}")
