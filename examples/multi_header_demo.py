"""多行表头演示"""
import sys
from pathlib import Path

# 添加项目路径到 sys.path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from openpyxl import Workbook

from finance_mask.scanner import ExcelScannerV2, find_header_row, build_header_name


def create_multi_header_workbook(filepath: Path):
    """创建包含多行表头的工作簿"""
    wb = Workbook()
    ws = wb.active
    ws.title = "财务报表"

    # 第一行：大类表头
    ws.cell(row=1, column=1, value="项目")
    ws.cell(row=1, column=2, value="本期")
    ws.cell(row=1, column=3, value="本期")
    ws.cell(row=1, column=4, value="上期")
    ws.cell(row=1, column=5, value="上期")

    # 第二行：小类表头
    ws.cell(row=2, column=1, value="")
    ws.cell(row=2, column=2, value="金额")
    ws.cell(row=2, column=3, value="占比")
    ws.cell(row=2, column=4, value="金额")
    ws.cell(row=2, column=5, value="占比")

    # 数据行
    data = [
        ("营业收入", 12345678.90, "100%", 11111111.11, "100%"),
        ("营业成本", 8765432.10, "71%", 7777777.77, "70%"),
        ("毛利润", 3580246.80, "29%", 3333333.34, "30%"),
        ("净利润", 2345678.90, "19%", 2222222.22, "20%"),
    ]

    for row_idx, (item, current_amount, current_ratio, prev_amount, prev_ratio) in enumerate(data, start=3):
        ws.cell(row=row_idx, column=1, value=item)
        ws.cell(row=row_idx, column=2, value=current_amount)
        ws.cell(row=row_idx, column=3, value=current_ratio)
        ws.cell(row=row_idx, column=4, value=prev_amount)
        ws.cell(row=row_idx, column=5, value=prev_ratio)

    wb.save(str(filepath))
    print(f"已创建多行表头工作簿: {filepath}")


def create_single_header_workbook(filepath: Path):
    """创建单行表头的工作簿"""
    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"

    # 单行表头
    headers = ["项目", "本期金额", "上期金额", "同比变动"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # 数据行
    data = [
        ("营业收入", 12345678.90, 11111111.11, "11.1%"),
        ("营业成本", 8765432.10, 7777777.77, "12.7%"),
        ("毛利润", 3580246.80, 3333333.34, "7.4%"),
        ("净利润", 2345678.90, 2222222.22, "5.6%"),
    ]

    for row_idx, (item, current, prev, change) in enumerate(data, start=2):
        ws.cell(row=row_idx, column=1, value=item)
        ws.cell(row=row_idx, column=2, value=current)
        ws.cell(row=row_idx, column=3, value=prev)
        ws.cell(row=row_idx, column=4, value=change)

    wb.save(str(filepath))
    print(f"已创建单行表头工作簿: {filepath}")


def demo_header_detection():
    """演示表头检测"""
    print("=" * 60)
    print("表头检测演示")
    print("=" * 60)
    print()

    # 创建临时目录
    import tempfile
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)

        # 创建不同类型的工作簿
        multi_header_file = tmp_path / "multi_header.xlsx"
        single_header_file = tmp_path / "single_header.xlsx"

        create_multi_header_workbook(multi_header_file)
        create_single_header_workbook(single_header_file)

        print()

        # 检测表头
        from openpyxl import load_workbook

        print("1. 多行表头检测:")
        wb = load_workbook(str(multi_header_file), read_only=True)
        ws = wb.active
        header_start, header_end = find_header_row(ws, candidates=None)
        print(f"   表头范围: 第 {header_start + 1} 行 到 第 {header_end + 1} 行")
        
        # 显示合并后的表头名称
        print("   合并后的表头名称:")
        for col_idx in range(ws.max_column or 0):
            combined_name = build_header_name(ws, col_idx, header_start, header_end)
            if combined_name:
                print(f"     列 {col_idx + 1}: {combined_name}")
        wb.close()

        print()
        print("2. 单行表头检测:")
        wb = load_workbook(str(single_header_file), read_only=True)
        ws = wb.active
        header_start, header_end = find_header_row(ws, candidates=None)
        print(f"   表头范围: 第 {header_start + 1} 行 到 第 {header_end + 1} 行")
        
        # 显示表头名称
        print("   表头名称:")
        for col_idx in range(ws.max_column or 0):
            combined_name = build_header_name(ws, col_idx, header_start, header_end)
            if combined_name:
                print(f"     列 {col_idx + 1}: {combined_name}")
        wb.close()


def demo_scanning():
    """演示扫描多行表头"""
    print()
    print("=" * 60)
    print("多行表头扫描演示")
    print("=" * 60)
    print()

    # 创建临时目录
    import tempfile
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)

        # 创建多行表头工作簿
        multi_header_file = tmp_path / "multi_header.xlsx"
        create_multi_header_workbook(multi_header_file)

        print()

        # 扫描
        scanner = ExcelScannerV2()
        sites = scanner.scan(multi_header_file)

        print(f"发现 {len(sites)} 个敏感位点:")
        for site in sites:
            print(f"  - {site.site_id}: {site.original_value} ({site.detected_type})")
            if site.location.column:
                print(f"    列名: {site.location.column}")


if __name__ == "__main__":
    demo_header_detection()
    demo_scanning()
