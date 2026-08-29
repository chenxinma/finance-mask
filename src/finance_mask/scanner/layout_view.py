"""
Excel 表格类型识别模块

使用 Rust 动态库识别 Excel 表是：
- Data 类型：表头+数据的一维表格
- Form 类型：key-value 排版的表单数据
"""
from __future__ import annotations

import ctypes
import json
import logging
import os
from ctypes import c_char_p
from enum import Enum
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


class SheetType(str, Enum):
    """工作表类型"""
    DATA = "Data"    # 表头+数据的一维表格
    FORM = "Form"    # key-value 排版的表单数据


class SheetClassification:
    """工作表分类结果"""
    
    def __init__(self, sheet_name: str, sheet_type: SheetType, confidence: float = 1.0):
        self.sheet_name = sheet_name
        self.sheet_type = sheet_type
        self.confidence = confidence
    
    def __repr__(self) -> str:
        return f"SheetClassification(sheet_name='{self.sheet_name}', sheet_type={self.sheet_type}, confidence={self.confidence})"


class LayoutViewClassifier:
    """Excel 表格类型分类器"""
    
    def __init__(self, lib_path: Optional[Path] = None):
        """
        初始化分类器
        
        Args:
            lib_path: Rust 动态库路径，默认在项目根目录的 lib 目录下
        """
        self._lib = None
        self._lib_path = lib_path or self._find_library()
    
    def _find_library(self) -> Path:
        """查找 Rust 动态库"""
        # 从当前文件向上查找项目根目录
        current_dir = Path(__file__).parent
        project_root = current_dir.parent.parent.parent
        
        # 根据操作系统确定库文件名
        if os.name == "posix":
            lib_name = "liblayout_view.so"
        elif os.name == "nt":
            lib_name = "layout_view.dll"
        else:
            raise OSError(f"不支持的操作系统: {os.name}")
        
        lib_path = project_root / "lib" / lib_name
        return lib_path
    
    def _load_library(self) -> None:
        """加载 Rust 动态库"""
        if self._lib is not None:
            return
        
        if not self._lib_path.exists():
            raise FileNotFoundError(f"Rust 动态库未找到: {self._lib_path}")
        
        try:
            self._lib = ctypes.CDLL(str(self._lib_path))
            
            # 定义函数签名
            self._lib.classify_excel_sheets_c.argtypes = [c_char_p]
            self._lib.classify_excel_sheets_c.restype = ctypes.POINTER(ctypes.c_char)
            
            self._lib.free_c_string.argtypes = [ctypes.POINTER(ctypes.c_char)]
            self._lib.free_c_string.restype = None
            
            logger.info(f"已加载 Rust 动态库: {self._lib_path}")
        except Exception as e:
            logger.error(f"加载 Rust 动态库失败: {e}")
            raise
    
    def classify(self, xlsx_path: Path) -> list[SheetClassification]:
        """
        分类 Excel 文件中的所有工作表
        
        Args:
            xlsx_path: Excel 文件路径
            
        Returns:
            工作表分类结果列表
        """
        self._load_library()
        
        # 转换为 C 字符串
        c_path = ctypes.c_char_p(str(xlsx_path).encode("utf-8"))
        
        # 调用 Rust 函数
        result_ptr = self._lib.classify_excel_sheets_c(c_path)
        
        if not result_ptr:
            logger.warning(f"Rust 函数返回空结果: {xlsx_path}")
            return []
        
        try:
            # 转换结果为 Python 字符串
            result_bytes = ctypes.cast(result_ptr, ctypes.c_char_p).value
            if result_bytes is None:
                return []
            
            result_str = result_bytes.decode("utf-8")
            
            # 解析 JSON 结果
            parsed_result = json.loads(result_str)
            
            # 转换为 SheetClassification 对象
            classifications = []
            for item in parsed_result:
                sheet_name = item.get("sheet_name", "")
                sheet_type_str = item.get("sheet_type", "Form")
                confidence = item.get("confidence", 1.0)
                
                # 转换为枚举类型
                try:
                    sheet_type = SheetType(sheet_type_str)
                except ValueError:
                    logger.warning(f"未知的工作表类型: {sheet_type_str}，默认为 Form")
                    sheet_type = SheetType.FORM
                
                classifications.append(SheetClassification(
                    sheet_name=sheet_name,
                    sheet_type=sheet_type,
                    confidence=confidence,
                ))
            
            return classifications
            
        except json.JSONDecodeError as e:
            logger.error(f"JSON 解析失败: {e}")
            return []
        except Exception as e:
            logger.error(f"处理分类结果时出错: {e}")
            raise
        finally:
            # 释放 Rust 分配的内存
            self._lib.free_c_string(result_ptr)
    
    def classify_with_fallback(self, xlsx_path: Path) -> list[SheetClassification]:
        """
        分类工作表，如果 Rust 库不可用则使用默认分类
        
        Args:
            xlsx_path: Excel 文件路径
            
        Returns:
            工作表分类结果列表
        """
        try:
            return self.classify(xlsx_path)
        except FileNotFoundError:
            logger.warning("Rust 动态库不可用，使用默认分类（所有工作表视为 Form）")
            return self._default_classification(xlsx_path)
        except Exception as e:
            logger.warning(f"Rust 分类失败: {e}，使用默认分类")
            return self._default_classification(xlsx_path)
    
    def _default_classification(self, xlsx_path: Path) -> list[SheetClassification]:
        """默认分类：所有工作表视为 Form 类型"""
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(xlsx_path), read_only=True)
            classifications = []
            for sheet_name in wb.sheetnames:
                classifications.append(SheetClassification(
                    sheet_name=sheet_name,
                    sheet_type=SheetType.FORM,
                    confidence=0.5,  # 低置信度表示这是默认分类
                ))
            wb.close()
            return classifications
        except Exception as e:
            logger.error(f"默认分类失败: {e}")
            return []
