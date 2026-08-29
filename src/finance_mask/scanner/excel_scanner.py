"""Excel 文件扫描实现"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..models.site import (
    Site,
    Location,
    SiteType,
    DetectedType,
    ActionType,
    DiscoveredBy,
)
from .base import BaseScanner
from .column_matcher import ColumnMatcher
from .patterns import PatternRegistry

logger = logging.getLogger(__name__)


class ExcelScanner(BaseScanner):
    """Excel 文件扫描器"""

    def __init__(
        self,
        column_matcher: Optional[ColumnMatcher] = None,
        pattern_registry: Optional[PatternRegistry] = None,
    ):
        self.column_matcher = column_matcher or ColumnMatcher()
        self.pattern_registry = pattern_registry or PatternRegistry()

    def scan(self, filepath: Path) -> list[Site]:
        """扫描 Excel 文件"""
        sites: list[Site] = []
        
        try:
            wb = load_workbook(str(filepath), read_only=True, data_only=True)
        except Exception as e:
            logger.error(f"无法打开 Excel 文件 {filepath}: {e}")
            return sites

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_sites = self._scan_sheet(filepath, sheet_name, ws)
            sites.extend(sheet_sites)

        wb.close()
        return sites

    def _scan_sheet(self, filepath: Path, sheet_name: str, ws) -> list[Site]:
        """扫描单个工作表"""
        sites: list[Site] = []
        header_row: dict[int, str] = {}  # 列索引 -> 列名
        header_found = False

        for row_idx, row in enumerate(ws.iter_rows(min_row=1), start=1):
            for col_idx, cell in enumerate(row, start=1):
                cell_value = cell.value
                if cell_value is None:
                    continue

                cell_str = str(cell_value).strip()
                if not cell_str:
                    continue

                cell_coord = f"{get_column_letter(col_idx)}{row_idx}"

                # 尝试识别表头行（前5行中查找）
                if row_idx <= 5 and not header_found:
                    rule = self.column_matcher.match(cell_str)
                    if rule:
                        header_row[col_idx] = cell_str
                        header_found = True
                        continue

                # 列头定位扫描
                if col_idx in header_row:
                    header_name = header_row[col_idx]
                    rule = self.column_matcher.match(header_name)
                    if rule:
                        site = self._create_site_from_column(
                            sheet_name=sheet_name,
                            cell_coord=cell_coord,
                            header_name=header_name,
                            value=cell_str,
                            rule=rule,
                        )
                        if site:
                            sites.append(site)
                            continue

                # 全文正则扫描（兜底）
                pattern_results = self.pattern_registry.scan_text(cell_str)
                for matched_value, detected_type, rule_name in pattern_results:
                    site = self._create_site_from_pattern(
                        sheet_name=sheet_name,
                        cell_coord=cell_coord,
                        value=matched_value,
                        detected_type=detected_type,
                    )
                    if site:
                        sites.append(site)

        return sites

    def _create_site_from_column(
        self,
        sheet_name: str,
        cell_coord: str,
        header_name: str,
        value: str,
        rule,
    ) -> Optional[Site]:
        """从列头规则创建位点"""
        # 根据规则类型确定默认动作
        if rule.action:
            action = rule.action
            params = rule.params
        else:
            action, params = self._get_default_action(rule.detected_type)

        location = Location(
            type=SiteType.EXCEL,
            sheet=sheet_name,
            cell=cell_coord,
            column=header_name,
        )

        site_id = self._generate_site_id("sheet", sheet_name, cell_coord)

        return Site(
            site_id=site_id,
            location=location,
            original_value=value,
            detected_type=rule.detected_type,
            discovered_by=DiscoveredBy.COLUMN_RULE,
            enabled=True,
            action=action,
            params=params,
        )

    def _create_site_from_pattern(
        self,
        sheet_name: str,
        cell_coord: str,
        value: str,
        detected_type: DetectedType,
    ) -> Optional[Site]:
        """从正则规则创建位点"""
        action, params = self._get_default_action(detected_type)

        location = Location(
            type=SiteType.EXCEL,
            sheet=sheet_name,
            cell=cell_coord,
        )

        site_id = self._generate_site_id("sheet", sheet_name, cell_coord)

        return Site(
            site_id=site_id,
            location=location,
            original_value=value,
            detected_type=detected_type,
            discovered_by=DiscoveredBy.FULLTEXT_SCAN,
            enabled=True,
            action=action,
            params=params,
        )

    def _get_default_action(self, detected_type: DetectedType) -> tuple[ActionType, dict]:
        """根据敏感类型获取默认脱敏动作"""
        if detected_type == DetectedType.AMOUNT:
            return ActionType.PRECISION, {"unit": "million", "decimal_places": 2}
        elif detected_type == DetectedType.ENTITY:
            return ActionType.ALIAS, {"prefix": "公司"}
        elif detected_type == DetectedType.PERSON:
            return ActionType.MASK_NAME, {}
        elif detected_type == DetectedType.ACCOUNT:
            return ActionType.MASK_ACCOUNT, {"keep_prefix": 3, "keep_suffix": 4}
        else:
            return ActionType.MASK, {}
