"""多行表头测试"""
import pytest
from pathlib import Path

from openpyxl import Workbook

from src.finance_mask.scanner import ExcelScannerV2
from src.finance_mask.scanner.header_finder import find_header_row, build_header_name


class TestMultiHeader:
    """多行表头测试"""

    @pytest.fixture
    def warehouse_workbook(self, tmp_path):
        """创建仓库入库单类型的工作簿（模拟实际文件）"""
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"

        # 第1-7行：标题和元数据区域
        ws.cell(row=2, column=1, value="仓库入库单")
        ws.cell(row=6, column=1, value="登记日期：")
        ws.cell(row=6, column=2, value="2026-04-23")

        # 第8行：大类表头
        ws.cell(row=8, column=1, value="单据信息")
        ws.cell(row=8, column=3, value="产品详情")
        ws.cell(row=8, column=5, value="入库详情")

        # 第9行：小类表头
        ws.cell(row=9, column=1, value="单据编号")
        ws.cell(row=9, column=2, value="单据日期")
        ws.cell(row=9, column=3, value="产品名称")
        ws.cell(row=9, column=4, value="产品规格")
        ws.cell(row=9, column=5, value="本次入库")
        ws.cell(row=9, column=6, value="补充说明")

        # 第10行开始：数据行
        data = [
            ("CK-A001", "2026-04-01", "纸箱", "50*40cm", 120, "日常补货"),
            ("CK-A002", "2026-04-01", "打包膜", "宽5cm", 350, "耗材领用"),
            ("CK-A003", "2026-04-02", "泡沫垫", "加厚款", 220, "新物料入库"),
            ("CK-A004", "2026-04-02", "胶带", "透明款1.5cm", 480, "批量采购"),
            ("CK-A005", "2026-04-03", "纸箱", "60*50cm", 90, "定制款补货"),
        ]

        for row_idx, (doc_no, doc_date, product_name, spec, quantity, note) in enumerate(data, start=10):
            ws.cell(row=row_idx, column=1, value=doc_no)
            ws.cell(row=row_idx, column=2, value=doc_date)
            ws.cell(row=row_idx, column=3, value=product_name)
            ws.cell(row=row_idx, column=4, value=spec)
            ws.cell(row=row_idx, column=5, value=quantity)
            ws.cell(row=row_idx, column=6, value=note)

        filepath = tmp_path / "warehouse.xlsx"
        wb.save(str(filepath))
        return filepath

    def test_find_multi_header(self, warehouse_workbook):
        """测试多行表头查找"""
        from openpyxl import load_workbook

        wb = load_workbook(str(warehouse_workbook), read_only=True)
        ws = wb.active

        header_start, header_end = find_header_row(ws, candidates=None)
        wb.close()

        # 应该找到第8-9行（索引7-8）
        assert header_start == 7
        assert header_end == 8

    def test_build_header_name(self, warehouse_workbook):
        """测试合并表头名称"""
        from openpyxl import load_workbook

        wb = load_workbook(str(warehouse_workbook), data_only=True)
        ws = wb.active

        # 构建合并后的表头名称
        header_names = []
        for col_idx in range(ws.max_column or 0):
            name = build_header_name(ws, col_idx, 7, 8)
            header_names.append(name)

        wb.close()

        # 验证表头名称（注意：第8行只有第1、3、5列有值）
        assert header_names[0] == "单据信息_单据编号"
        assert header_names[1] == "单据日期"  # 第8行第2列为空
        assert header_names[2] == "产品详情_产品名称"
        assert header_names[3] == "产品规格"  # 第8行第4列为空
        assert header_names[4] == "入库详情_本次入库"
        assert header_names[5] == "补充说明"  # 第8行第6列为空

    def test_scan_multi_header(self, warehouse_workbook):
        """测试扫描多行表头工作表"""
        scanner = ExcelScannerV2()
        sites = scanner.scan(warehouse_workbook)

        # 应该发现敏感位点
        assert len(sites) > 0
        print(f"发现 {len(sites)} 个敏感位点")

        # 检查是否包含账号位点（单据编号）
        account_sites = [s for s in sites if s.detected_type.value == "account"]
        assert len(account_sites) > 0
        print(f"账号位点: {len(account_sites)} 个")

        # 检查是否包含金额位点（本次入库）
        amount_sites = [s for s in sites if s.detected_type.value == "amount" and "入库" in (s.location.column or "")]
        assert len(amount_sites) > 0
        print(f"入库金额位点: {len(amount_sites)} 个")

    def test_scan_actual_file(self):
        """测试扫描实际文件（如果存在）"""
        actual_file = Path("examples/仓库入库1.xlsx")
        if not actual_file.exists():
            pytest.skip("实际文件不存在")

        scanner = ExcelScannerV2()
        sites = scanner.scan(actual_file)

        # 应该发现敏感位点
        assert len(sites) > 0
        print(f"实际文件发现 {len(sites)} 个敏感位点")

        # 检查列名是否正确合并
        column_names = set()
        for site in sites:
            if site.location.column:
                column_names.add(site.location.column)

        print(f"列名: {column_names}")

        # 验证列名包含合并后的格式
        assert any("_" in name for name in column_names)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
