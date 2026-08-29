from openpyxl import load_workbook
from openpyxl.styles.proxy import StyleProxy
from openpyxl.worksheet.worksheet import Worksheet

def improved_find_header_row(
    sheet,
    header_candidates=None,
    max_scan_rows=10,
    blanks_ratio=0.5,
    fuzzy_match_threshold=0.8,
):
    """
    改进的表头行查找算法
    新增特征：文本属性（中文/名词）、数字占比、与数据行的区分度、唯一值密度
    """
    # 预处理：获取所有扫描行的原始值和处理后的值（去空、去空格）
    scan_rows = []
    max_row = min(max_scan_rows, sheet.max_row)
    ncols = max(sheet.max_column, 1) if sheet.max_column else 1  # 处理空表格
    max_row_value_count = 0
    for r_idx in range(1, max_row + 1):
        row_cells = sheet[r_idx]  # 按行索引取单元格（从1开始）
        row_values = [cell.value for cell in row_cells]
        # 处理后的值：非空、去空格、转字符串
        processed_vals = [
            str(v).strip()
            for v in row_values
            if v is not None and str(v).strip() != ""
        ]
        scan_rows.append(
            {
                "r_idx": r_idx - 1,  # 原始返回的是0开始的索引
                "raw_values": row_values,
                "processed": processed_vals,
                "non_blank_count": len(processed_vals),
            }
        )
        if len(row_values) > max_row_value_count:
            max_row_value_count = len(row_values)

    # 不对单列数据进行处理
    if max_row_value_count < 2:
        return -1

    # 遍历每一行，计算综合评分
    header_scores = []
    for row_data in scan_rows:
        r_idx = row_data["r_idx"]
        raw_vals = row_data["raw_values"]
        processed_vals = row_data["processed"]
        non_blank_count = row_data["non_blank_count"]
        score = 0.0

        # 规则1：空值比例过滤（不满足则直接跳过，评分为0）
        blank_ratio = 1 - (non_blank_count / ncols) if ncols > 0 else 1.0
        if blank_ratio > blanks_ratio:
            header_scores.append((r_idx, 0.0))
            continue
        score += 1.0  # 满足空值比例，基础分

        # 规则2：唯一值比例（表头通常唯一值多）
        unique_vals = set(processed_vals)
        unique_ratio = (
            len(unique_vals) / len(processed_vals) if processed_vals else 0.0
        )
        score += unique_ratio * 1.0  # 权重1.0

        # 规则3：数字占比（表头通常数字少，数据行数字多）
        num_count = 0
        for val in processed_vals:
            # 判断是否为纯数字（整数/小数，排除身份证等长数字字符串）
            if (
                re.match(r"^-?\d+(\.\d+)?$", val) and len(val) <= 15
            ):  # 15位以内纯数字视为数值
                num_count += 1
        num_ratio = num_count / len(processed_vals) if processed_vals else 1.0
        score += (1 - num_ratio) * 1.5  # 数字占比越低，得分越高，权重1.5

        # 规则4：中文字符占比（表头通常包含中文，数据行可能少）
        chinese_count = 0
        for val in processed_vals:
            chinese_count += sum(1 for c in val if "\u4e00" <= c <= "\u9fff")
        chinese_ratio = (
            chinese_count / sum(len(val) for val in processed_vals)
            if processed_vals
            else 0.0
        )
        score += chinese_ratio * 2.0  # 权重2.0（中文对表头识别更重要）

        # 规则5：下方行的数据验证（表头下方应该是数据行，满足数据特征）
        # 取当前行下一行（如果存在），判断是否为数据行（数字占比高、非空）
        next_row_idx = r_idx + 1
        if next_row_idx < len(scan_rows):
            next_row = scan_rows[next_row_idx]
            next_processed = next_row["processed"]
            next_non_blank = next_row["non_blank_count"]
            if next_non_blank / ncols >= (1 - blanks_ratio):  # 下一行非空比例足够
                # 计算下一行的数字占比
                next_num_count = 0
                for val in next_processed:
                    if re.match(r"^-?\d+(\.\d+)?$", val) and len(val) <= 15:
                        next_num_count += 1
                next_num_ratio = (
                    next_num_count / len(next_processed) if next_processed else 0.0
                )
                if next_num_ratio > 0.5:  # 下一行数字占比超过50%，视为数据行
                    score += 1.5  # 权重1.5

        # 规则6：关键字模糊匹配（如果有候选关键字）
        if header_candidates:
            matched_count = 0
            for candidate in header_candidates:
                candidate_lower = candidate.lower().strip()
                # 模糊匹配：候选词在单元格值中（忽略大小写）
                if any(
                    candidate_lower in str(v).lower().strip()
                    for v in raw_vals
                    if v is not None
                ):
                    matched_count += 1
            match_ratio = (
                matched_count / len(header_candidates) if header_candidates else 0.0
            )
            if match_ratio >= fuzzy_match_threshold:
                score += match_ratio * 3.0  # 权重3.0（关键字匹配优先级最高）
            else:
                score *= 0.5  # 匹配不足，降低评分

        header_scores.append((r_idx, score))

    # 找到评分最高的行（如果多个最高分，取第一个）
    if not header_scores:
        return -1
    header_scores.sort(key=lambda x: (-x[1], x[0]))  # 按评分降序、索引升序
    best_r_idx, best_score = header_scores[0]

    # 如果最高分是0，说明没有符合条件的行
    return best_r_idx if best_score > 0 else -1
    
# Usage
# import pandas as pd
# header_row = improved_find_header_row(
#     sheet, header_candidates=None
# )
# if header_row == -1:
#     log.error("sheet no header found")
#     return
# 
# fname = f"{Path(filename).stem}_{sheet_name}"
# # 读取数据
# df = pd.read_excel(
#     original_path, sheet_name=sheet_name, header=header_row
# )