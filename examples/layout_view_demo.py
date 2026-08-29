"""Excel 表格类型识别演示"""
import sys
from pathlib import Path

# 添加项目路径到 sys.path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

from openpyxl import Workbook

from finance_mask.scanner import ExcelScannerV2, LayoutViewClassifier, SheetType


def create_data_workbook(filepath: Path):
    """创建 Data 类型的工作簿（表头+数据的一维表格）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"

    # 表头行
    headers = ["项目", "本期金额", "上期金额", "同比变动"]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    # 数据行
    data = [
        ("营业收入", 12345678.90, 11111111.11, "11.1%"),
        ("营业成本", 8765432.10, 7777777.77, "12.7%"),
        ("毛利润", 3580246.80, 3333333.34, "7.4%"),
        ("净利润", 2345678.90, 2222222.22, "5.6%"),
    ]

    for row_idx, (item, current, prev, change) in enumerate(data, start=2):
        ws.cell(row=row_idx, column=1, value=item)
        ws.cell(row=row_idx, column=2, value=current)
        ws.cell(row=row_idx, column=3, value=prev)
        ws.cell(row=row_idx, column=4, value=change)

    wb.save(str(filepath))
    print(f"已创建 Data 类型工作簿: {filepath}")


def create_form_workbook(filepath: Path):
    """创建 Form 类型的工作簿（key-value 排版的表单数据）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "基本信息"

    # 表单数据（key-value 格式）
    form_data = [
        ("公司名称", "阿里巴巴集团"),
        ("统一社会信用代码", "91330000799210495T"),
        ("法定代表人", "张三"),
        ("注册资本", "1000000万元"),
        ("成立日期", "2020-01-01"),
        ("营业收入", "1234567890.12元"),
        ("净利润", "234567890.50元"),
    ]

    for row_idx, (key, value) in enumerate(form_data, start=1):
        ws.cell(row=row_idx, column=1, value=key)
        ws.cell(row=row_idx, column=2, value=value)

    wb.save(str(filepath))
    print(f"已创建 Form 类型工作簿: {filepath}")


def create_mixed_workbook(filepath: Path):
    """创建包含多种类型工作表的工作簿"""
    wb = Workbook()

    # Data 类型工作表
    ws1 = wb.active
    ws1.title = "利润表"
    headers = ["项目", "金额", "占比"]
    for col_idx, header in enumerate(headers, start=1):
        ws1.cell(row=1, column=col_idx, value=header)
    data = [
        ("营业收入", 12345678.90, "100%"),
        ("营业成本", 8765432.10, "71%"),
        ("毛利润", 3580246.80, "29%"),
    ]
    for row_idx, (item, amount, ratio) in enumerate(data, start=2):
        ws1.cell(row=row_idx, column=1, value=item)
        ws1.cell(row=row_idx, column=2, value=amount)
        ws1.cell(row=row_idx, column=3, value=ratio)

    # Form 类型工作表
    ws2 = wb.create_sheet("公司信息")
    form_data = [
        ("公司名称", "腾讯科技有限公司"),
        ("法定代表人", "李四"),
        ("注册资本", "500000万元"),
    ]
    for row_idx, (key, value) in enumerate(form_data, start=1):
        ws2.cell(row=row_idx, column=1, value=key)
        ws2.cell(row=row_idx, column=2, value=value)

    wb.save(str(filepath))
    print(f"已创建混合类型工作簿: {filepath}")


def demo_classification():
    """演示工作表类型分类"""
    print("=" * 60)
    print("工作表类型分类演示")
    print("=" * 60)
    print()

    # 创建临时目录
    import tempfile
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)

        # 创建不同类型的工作簿
        data_file = tmp_path / "data_report.xlsx"
        form_file = tmp_path / "form_report.xlsx"
        mixed_file = tmp_path / "mixed_report.xlsx"

        create_data_workbook(data_file)
        create_form_workbook(form_file)
        create_mixed_workbook(mixed_file)

        print()

        # 分类演示
        classifier = LayoutViewClassifier()

        print("1. Data 类型工作簿分类结果:")
        classifications = classifier.classify_with_fallback(data_file)
        for c in classifications:
            print(f"   - {c.sheet_name}: {c.sheet_type} (置信度: {c.confidence:.2f})")

        print()
        print("2. Form 类型工作簿分类结果:")
        classifications = classifier.classify_with_fallback(form_file)
        for c in classifications:
            print(f"   - {c.sheet_name}: {c.sheet_type} (置信度: {c.confidence:.2f})")

        print()
        print("3. 混合类型工作簿分类结果:")
        classifications = classifier.classify_with_fallback(mixed_file)
        for c in classifications:
            print(f"   - {c.sheet_name}: {c.sheet_type} (置信度: {c.confidence:.2f})")


def demo_scanning():
    """演示不同类型工作表的扫描"""
    print()
    print("=" * 60)
    print("不同类型工作表扫描演示")
    print("=" * 60)
    print()

    # 创建临时目录
    import tempfile
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir)

        # 创建不同类型的工作簿
        data_file = tmp_path / "data_report.xlsx"
        form_file = tmp_path / "form_report.xlsx"

        create_data_workbook(data_file)
        create_form_workbook(form_file)

        print()

        # 扫描演示
        scanner = ExcelScannerV2()

        print("1. 扫描 Data 类型工作簿:")
        sites = scanner.scan(data_file)
        print(f"   发现 {len(sites)} 个敏感位点")
        for site in sites[:3]:  # 只显示前3个
            print(f"   - {site.site_id}: {site.original_value} ({site.detected_type})")
        if len(sites) > 3:
            print(f"   ... 还有 {len(sites) - 3} 个位点")

        print()
        print("2. 扫描 Form 类型工作簿:")
        sites = scanner.scan(form_file)
        print(f"   发现 {len(sites)} 个敏感位点")
        for site in sites[:3]:  # 只显示前3个
            print(f"   - {site.site_id}: {site.original_value} ({site.detected_type})")
        if len(sites) > 3:
            print(f"   ... 还有 {len(sites) - 3} 个位点")


if __name__ == "__main__":
    demo_classification()
    demo_scanning()
