"""创建示例财务文件"""
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def create_sample_excel():
    """创建示例 Excel 文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"

    # 表头
    ws["A1"] = "项目"
    ws["B1"] = "本期金额（元）"
    ws["C1"] = "上期金额（元）"
    ws["D1"] = "同比变动"

    # 数据
    data = [
        ("营业收入", 1234567890.12, 1111111111.11, "11.1%"),
        ("营业成本", 876543210.00, 777777777.77, "12.7%"),
        ("毛利润", 358024680.12, 333333333.34, "7.4%"),
        ("净利润", 234567890.50, 222222222.22, "5.6%"),
        ("客户名称", "阿里巴巴集团", "腾讯科技有限公司", "-"),
        ("联系人", "张三", "李四", "-"),
        ("合同编号", "HT-2024-001234", "HT-2023-005678", "-"),
    ]

    for i, (item, current, prev, change) in enumerate(data, start=2):
        ws[f"A{i}"] = item
        ws[f"B{i}"] = current
        ws[f"C{i}"] = prev
        ws[f"D{i}"] = change

    wb.save("examples/sample_report.xlsx")
    print("已创建示例 Excel: examples/sample_report.xlsx")


def create_sample_pptx():
    """创建示例 PPT 文件"""
    prs = Presentation()
    
    # 幻灯片1：标题页
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "2024年度财务报告"
    slide1.placeholders[1].text = "机密文件 - 仅限内部使用"

    # 幻灯片2：关键指标
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    slide2.shapes.title.text = "关键财务指标"
    
    left = Inches(1)
    top = Inches(2)
    width = Inches(8)
    height = Inches(3)
    
    txBox = slide2.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.text = "2024年度营业收入为 1,234,567,890.12 元，同比增长 11.1%\n净利润为 234,567,890.50 元，同比增长 5.6%"

    # 幻灯片3：表格
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    
    rows = 5
    cols = 4
    left = Inches(1)
    top = Inches(1)
    width = Inches(8)
    height = Inches(4)
    
    table_shape = slide3.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    
    # 表头
    headers = ["项目", "金额（万元）", "占比", "备注"]
    for i, header in enumerate(headers):
        table.cell(0, i).text = header
    
    # 数据
    table_data = [
        ("营业收入", "123,456.79", "100%", "主要收入来源"),
        ("营业成本", "87,654.32", "71.0%", "成本控制良好"),
        ("毛利润", "35,802.47", "29.0%", "毛利率稳定"),
        ("净利润", "23,456.79", "19.0%", "盈利能力强"),
    ]
    
    for i, (item, amount, ratio, note) in enumerate(table_data, start=1):
        table.cell(i, 0).text = item
        table.cell(i, 1).text = amount
        table.cell(i, 2).text = ratio
        table.cell(i, 3).text = note

    prs.save("examples/sample_report.pptx")
    print("已创建示例 PPT: examples/sample_report.pptx")


if __name__ == "__main__":
    create_sample_excel()
    create_sample_pptx()
    print("\n示例文件创建完成！")
    print("\n使用方法：")
    print("1. 生成策略: python -m finance_mask generate -i examples/sample_report.xlsx -o examples/策略.yaml")
    print("2. 执行脱敏: python -m finance_mask redact -i examples/sample_report.xlsx -s examples/策略.yaml -o examples/输出/")
