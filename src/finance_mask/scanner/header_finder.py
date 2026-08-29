"""
表头行查找算法 - 改进版本

基于 excel_convert 项目的 header_finder.py 优化

主要改进：
1. 支持多行表头：返回 (start_row, end_row)
2. 新增注释行识别：避免误判为表头
3. 新增数据行识别：更准确地识别数据行
4. 使用 Shannon 熵计算内容多样性
5. 支持合并单元格
6. 支持多行表头名称合并
"""
from __future__ import annotations

import math
import re
from typing import Optional, Tuple, List

from openpyxl.worksheet.worksheet import Worksheet

MAX_HEADER_ROWS = 5


def is_annotation_row(row_values: tuple) -> bool:
    """
    检查是否为注释/说明行

    返回 True 如果：
    - 第一个非空值以 '注释', 'Notes', 'Description', '注释：' 开头
    - 或者 50%+ 的非空值长度小于 5 个字符（排除中文字符）
    """
    non_empty = [str(v) for v in row_values if v is not None and str(v).strip()]

    if not non_empty:
        return False

    # 检查第一个非空值是否以注释关键字开头
    first_value = non_empty[0]
    annotation_prefixes = ("注释", "Notes", "Description", "注释：")
    if any(first_value.startswith(prefix) for prefix in annotation_prefixes):
        return True

    # 检查 50%+ 的非空值是否长度小于 5 个字符
    # 排除中文字符（通常 1-2 个字符但有意义）
    chinese_pattern = re.compile(r"[\u4e00-\u9fff]")
    short_count = 0
    for v in non_empty:
        # 只有不含中文且长度 < 5 才算短
        if not chinese_pattern.search(v) and len(v) < 5:
            short_count += 1
    return short_count / len(non_empty) >= 0.5


def is_data_row(row_values: tuple) -> bool:
    """
    检测是否为数据行（对比表头行）

    返回 True 如果：
    - 数字比例 > 60%
    - 或唯一值比例 < 40% 且不像表头行
    - 或第一个值是数字（可能是 ID/序号）且内容混合
    - 或短值且高唯一性（典型数据模式）
    """
    non_empty = [str(v) for v in row_values if v is not None]

    if not non_empty:
        return False

    # 统计字符类型
    chinese_count = 0
    total_chars = 0
    for val in non_empty:
        for char in val:
            total_chars += 1
            if "\u4e00" <= char <= "\u9fff":
                chinese_count += 1

    chinese_ratio = chinese_count / total_chars if total_chars > 0 else 0

    # 数字比例 > 60%
    number_pattern = re.compile(r"^-?\d+(\.\d+)?$")
    number_count = 0
    for val in non_empty:
        if number_pattern.match(val) and len(val.replace("-", "").replace(".", "")) <= 15:
            number_count += 1

    number_ratio = number_count / len(non_empty)
    if number_ratio > 0.6:
        return True

    # 唯一值比例 < 40%
    unique_count = len(set(non_empty))
    unique_ratio = unique_count / len(non_empty)
    if unique_ratio < 0.4:
        # 如果行有很多中文字符且唯一值比例低，可能是重复列名的表头行
        if chinese_ratio > 0.5 and len(non_empty) > 10:
            return False
        return True

    # 检查第一个值是否为数字（可能是 ID/序号）
    if len(non_empty) >= 3 and number_pattern.match(non_empty[0]):
        first_val = non_empty[0]
        # 如果第一个值是合理的 ID 长度（不像身份证号那么长）
        if 1 <= len(first_val) <= 8:
            return True

    # 检查混合内容：一些数字 + 中文文本（典型数据行模式）
    if number_count > 0 and chinese_count > 0 and len(non_empty) >= 3:
        return True

    # 检查典型数据模式：短值、高唯一性、字母数字内容
    if len(non_empty) >= 3 and chinese_ratio > 0:
        avg_len = sum(len(v) for v in non_empty) / len(non_empty)
        if avg_len < 15 and unique_ratio > 0.8:
            # 检查是否有字母数字 ID（如 'E0001', 'ABC123'）
            alphanumeric_pattern = re.compile(r"^[A-Za-z0-9]+$")
            alpha_count = sum(
                1 for v in non_empty
                if alphanumeric_pattern.match(v) and not number_pattern.match(v)
            )
            if alpha_count > 0:
                return True

    return False


def get_vertical_value(sheet: Worksheet, row: int, col: int) -> str:
    """
    获取单元格值，处理合并单元格

    Args:
        sheet: openpyxl 工作表对象
        row: 0-based 行索引
        col: 0-based 列索引

    Returns:
        单元格值的字符串，如果没有值返回 ""
    """
    # 转换为 1-based 索引（openpyxl 使用 1-based）
    row_1based = row + 1
    col_1based = col + 1

    cell = sheet.cell(row=row_1based, column=col_1based)
    if cell.value is not None:
        return str(cell.value)

    # 检查合并范围（只读模式下可能不可用）
    try:
        for merged_range in sheet.merged_cells:
            if cell.coordinate in merged_range:
                # 获取主单元格（合并范围的左上角）
                master_cell = sheet.cell(
                    row=merged_range.min_row, column=merged_range.min_col
                )
                if master_cell.value is not None:
                    return str(master_cell.value)
    except AttributeError:
        # 只读模式下 merged_cells 可能不可用
        pass

    return ""


def build_header_name(
    sheet: Worksheet, col_idx: int, start_row: int, end_row: int
) -> str:
    """
    通过合并多行的值构建表头名称

    对于 start_row 到 end_row（0-based，包含）的每一行，
    通过 get_vertical_value 获取值，过滤空字符串，然后用 "_" 连接。

    Args:
        sheet: openpyxl 工作表对象
        col_idx: 0-based 列索引
        start_row: 0-based 起始行（包含）
        end_row: 0-based 结束行（包含）

    Returns:
        合并后的表头名称，非空值用 "_" 连接
    """
    parts = []
    for row in range(start_row, end_row + 1):
        value = get_vertical_value(sheet, row, col_idx)
        if value and value.strip():
            cleaned = re.sub(r"[\n /]+", "_", value.strip())
            parts.append(cleaned)
    return "_".join(parts)


def find_header_row(
    sheet: Worksheet, candidates: Optional[List[str]] = None
) -> Tuple[int, int]:
    """
    查找 Excel 工作表中的表头行范围

    扫描前 20 行，对每行评分，然后选择最佳表头行并向上向下扩展以捕获多行表头。

    Args:
        sheet: openpyxl 工作表对象
        candidates: 可选的候选列名列表，用于关键字匹配

    Returns:
        (start_row, end_row) 0-based 包含索引，如果未找到返回 (-1, -1)
    """
    max_scan_rows = 20
    fuzzy_match_threshold = 0.8

    row_info = []

    for row_idx in range(1, min(max_scan_rows + 1, sheet.max_row + 1)):
        row = list(sheet.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
        if not row:
            continue

        row_values = row[0]
        non_empty_values = [
            cell for cell in row_values if cell is not None and str(cell).strip()
        ]

        score = _calculate_row_score(
            row_values, row_idx, sheet, candidates, fuzzy_match_threshold
        )
        if score > 0:
            row_info.append({
                "row_0based": row_idx - 1,
                "score": score,
                "non_empty_count": len(non_empty_values),
                "is_annotation": is_annotation_row(tuple(row_values)),
                "is_data": is_data_row(row_values),
                "values": row_values,
            })

    if not row_info:
        return (-1, -1)

    # 按分数降序排序，然后按非空单元格数降序排序
    row_info.sort(key=lambda x: (-x["score"], -x["non_empty_count"]))

    # 选择最佳表头行：最高分，非注释行，优先选择内容更多的行
    best_row_info = None
    for info in row_info:
        if info["is_annotation"]:
            continue
        best_row_info = info
        break

    if best_row_info is None:
        return (-1, -1)

    best_row = best_row_info["row_0based"]
    best_non_empty_count = best_row_info["non_empty_count"]

    # 向上扩展以捕获主表头上方的标题/元数据行
    header_start = best_row
    for offset in range(1, MAX_HEADER_ROWS):
        prev_row_idx = best_row - offset
        if prev_row_idx < 0:
            break

        prev_row = list(
            sheet.iter_rows(
                min_row=prev_row_idx + 1, max_row=prev_row_idx + 1, values_only=True
            )
        )
        if not prev_row:
            break

        prev_values = tuple(prev_row[0])
        non_empty = [v for v in prev_values if v is not None and str(v).strip()]

        # 空行停止
        if len(non_empty) == 0:
            break

        # 如果前一行是数据行，停止
        if is_data_row(prev_values):
            break

        # 如果前一行的非空单元格数显著少于最佳行，停止
        # 这防止稀疏的元数据行被包含
        if len(non_empty) < max(3, best_non_empty_count * 0.3):
            break

        header_start = prev_row_idx

    # 向下扩展以形成多行表头
    header_rows = [header_start]
    for offset in range(header_start + 1, min(header_start + MAX_HEADER_ROWS, sheet.max_row)):
        next_row = list(
            sheet.iter_rows(min_row=offset + 1, max_row=offset + 1, values_only=True)
        )
        if not next_row:
            break

        next_values = tuple(next_row[0])
        total_cols = len(next_values)
        non_empty = [v for v in next_values if v is not None and str(v).strip()]
        non_empty_count = len(non_empty)

        # 在非常稀疏的行停止
        if total_cols > 0:
            blanks_ratio = (total_cols - non_empty_count) / total_cols
            if blanks_ratio > 0.8 and non_empty_count < 5:
                break

        # 如果遇到数据行，停止
        if is_data_row(next_values):
            break

        header_rows.append(offset)

    return (header_rows[0], header_rows[-1])


def _calculate_row_score(
    row: tuple,
    row_idx: int,
    sheet: Worksheet,
    candidates: Optional[List[str]],
    fuzzy_match_threshold: float,
) -> float:
    """计算行的多维度评分"""
    score = 0.0

    non_empty_values = [str(cell) for cell in row if cell is not None]
    if not non_empty_values:
        return 0.0

    non_empty_count = len(non_empty_values)

    score += _score_unique_ratio(non_empty_values)
    score += _score_number_ratio(non_empty_values)
    score += _score_chinese_ratio(non_empty_values)
    score += _score_content_diversity(non_empty_values)
    score += _score_data_validation(row_idx, sheet)

    # 对非空单元格更多的行加分（更可能是表头行）
    if non_empty_count >= 5:
        score += 0.5
    if non_empty_count >= 10:
        score += 0.5

    # 对非常稀疏的行扣分（可能是元数据，不是表头）
    if non_empty_count <= 2:
        score *= 0.5

    # 对典型表头行模式加分：
    # - 全是中文字符（列名通常是中文）
    # - 短值（列名通常很短）
    chinese_pattern = re.compile(r"^[\u4e00-\u9fff\s/%()（）]+$")
    all_chinese = all(chinese_pattern.match(v) for v in non_empty_values)
    avg_len = sum(len(v) for v in non_empty_values) / non_empty_count if non_empty_count > 0 else 0

    if all_chinese and avg_len <= 10:
        score += 1.0  # 纯中文短列名强加分

    if candidates:
        score += _score_keyword_match(
            non_empty_values, candidates, fuzzy_match_threshold
        )

    return score


def _score_content_diversity(values: List[str]) -> float:
    """
    规则 7：字符类型多样性的 Shannon 熵（权重 2.0）

    使用 Shannon 熵衡量字符类型多样性：
    - C: 中文字符
    - L: 英文字母
    - N: 数字
    - O: 其他

    表头行通常混合中文字段名和英文缩写，熵较高。
    纯数据行（全是数字或全是文本）熵较低。
    """
    def char_type(c: str) -> str:
        if "\u4e00" <= c <= "\u9fff":
            return "C"
        if c.isalpha() and c.isascii():
            return "L"
        if c.isdigit():
            return "N"
        return "O"

    if not values:
        return 0.0

    # 统计所有单元格的字符类型
    type_counts = {"C": 0, "L": 0, "N": 0, "O": 0}
    total = 0
    for val in values:
        for c in str(val):
            type_counts[char_type(c)] += 1
            total += 1

    if total == 0:
        return 0.0

    # Shannon 熵：H = -Σ p_i * ln(p_i)
    entropy = 0.0
    for t in type_counts:
        p = type_counts[t] / total
        if p > 0:
            entropy -= p * math.log(p)

    # 4 种类型的最大熵是 ln(4) ≈ 1.386
    # 归一化到 [0, 1] 并按权重缩放
    max_entropy = math.log(4)
    normalized = entropy / max_entropy
    return normalized * 2.0


def _score_unique_ratio(values: List[str]) -> float:
    """规则 2：唯一值比例（权重 1.0）"""
    unique_count = len(set(values))
    return (unique_count / len(values)) * 1.0


def _score_number_ratio(values: List[str]) -> float:
    """规则 3：低数字比例（权重 1.5）"""
    number_pattern = re.compile(r"^-?\d+(\.\d+)?$")
    number_count = 0

    for val in values:
        if number_pattern.match(val) and len(val.replace("-", "").replace(".", "")) <= 15:
            number_count += 1

    ratio = number_count / len(values)
    return (1 - ratio) * 1.5


def _score_chinese_ratio(values: List[str]) -> float:
    """规则 4：中文字符比例（权重 2.0）"""
    total_chars = 0
    chinese_chars = 0

    for val in values:
        for char in str(val):
            total_chars += 1
            if "\u4e00" <= char <= "\u9fff":
                chinese_chars += 1

    if total_chars == 0:
        return 0.0

    return (chinese_chars / total_chars) * 2.0


def _score_data_validation(row_idx: int, sheet: Worksheet) -> float:
    """规则 5：数据验证 - 下一行应该是数据（权重 1.0）"""
    if row_idx >= sheet.max_row:
        return 0.0

    next_row = list(
        sheet.iter_rows(min_row=row_idx + 1, max_row=row_idx + 1, values_only=True)
    )
    if not next_row:
        return 0.0

    next_values = [cell for cell in next_row[0] if cell is not None]
    if not next_values:
        return 0.0

    number_pattern = re.compile(r"^-?\d+(\.\d+)?$")
    number_count = 0

    for val in next_values:
        val_str = str(val)
        if number_pattern.match(val_str) and len(val_str.replace("-", "").replace(".", "")) <= 15:
            number_count += 1

    if number_count / len(next_values) > 0.5:
        return 1.0

    return 0.0


def _score_keyword_match(
    values: List[str], candidates: List[str], threshold: float
) -> float:
    """规则 6：关键字匹配（权重 3.0）"""
    if not candidates:
        return 0.0

    match_count = 0
    for val in values:
        for candidate in candidates:
            if candidate.lower() in val.lower() or val.lower() in candidate.lower():
                match_count += 1
                break

    match_rate = match_count / len(values)

    if match_rate >= threshold:
        return match_rate * 3.0
    else:
        return (match_rate * 3.0) / 2


# 向后兼容的接口
def find_header_row_single(sheet: Worksheet, candidates: Optional[List[str]] = None) -> int:
    """
    查找单行表头（向后兼容接口）

    Args:
        sheet: openpyxl 工作表对象
        candidates: 可选的候选列名列表

    Returns:
        表头行索引（0-based），如果未找到返回 -1
    """
    start, end = find_header_row(sheet, candidates)
    return start


def find_header_row_from_dataframe(
    df_headers: list[str],
    header_candidates: Optional[list[str]] = None,
    fuzzy_match_threshold: float = 0.8,
) -> bool:
    """
    验证 DataFrame 的列名是否为有效的表头

    Args:
        df_headers: DataFrame 的列名列表
        header_candidates: 候选表头关键字列表
        fuzzy_match_threshold: 模糊匹配阈值

    Returns:
        是否为有效表头
    """
    if not df_headers:
        return False

    # 检查是否包含空列名
    non_empty_headers = [h for h in df_headers if h and str(h).strip()]
    if len(non_empty_headers) < len(df_headers) * 0.5:
        return False

    # 如果有候选关键字，检查匹配度
    if header_candidates:
        matched_count = 0
        for candidate in header_candidates:
            candidate_lower = candidate.lower().strip()
            if any(candidate_lower in str(h).lower().strip() for h in df_headers):
                matched_count += 1
        match_ratio = matched_count / len(header_candidates)
        return match_ratio >= fuzzy_match_threshold

    # 检查是否包含典型的表头特征
    header_features = [
        any("\u4e00" <= c <= "\u9fff" for c in str(h))  # 包含中文
        for h in non_empty_headers
    ]
    chinese_ratio = sum(header_features) / len(header_features) if header_features else 0.0

    # 如果超过 50% 的列名包含中文，认为是有效表头
    return chinese_ratio > 0.5
